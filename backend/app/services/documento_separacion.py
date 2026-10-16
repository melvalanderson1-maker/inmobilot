"""Genera el contrato / proforma de separación a partir de la plantilla Word."""
import io
import os
import shutil
import subprocess
import tempfile
import uuid
from decimal import Decimal

import docx

_BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # backend/app
RUTA_PLANTILLA = os.path.join(_BASE_DIR, "templates", "plantilla_separacion.docx")
DIR_DOCUMENTOS = os.path.join(_BASE_DIR, "static", "separaciones")
os.makedirs(DIR_DOCUMENTOS, exist_ok=True)

_UNIDADES = ["", "UN", "DOS", "TRES", "CUATRO", "CINCO", "SEIS", "SIETE", "OCHO", "NUEVE", "DIEZ", "ONCE",
             "DOCE", "TRECE", "CATORCE", "QUINCE", "DIECISEIS", "DIECISIETE", "DIECIOCHO", "DIECINUEVE",
             "VEINTE", "VEINTIUN", "VEINTIDOS", "VEINTITRES", "VEINTICUATRO", "VEINTICINCO", "VEINTISEIS",
             "VEINTISIETE", "VEINTIOCHO", "VEINTINUEVE"]
_DECENAS = ["", "", "", "TREINTA", "CUARENTA", "CINCUENTA", "SESENTA", "SETENTA", "OCHENTA", "NOVENTA"]
_CENTENAS = ["", "CIENTO", "DOSCIENTOS", "TRESCIENTOS", "CUATROCIENTOS", "QUINIENTOS", "SEISCIENTOS",
             "SETECIENTOS", "OCHOCIENTOS", "NOVECIENTOS"]


def _menor_mil(n: int) -> str:
    if n == 100:
        return "CIEN"
    c, r = divmod(n, 100)
    partes = [_CENTENAS[c]] if c else []
    if r < 30:
        if r:
            partes.append(_UNIDADES[r])
    else:
        d, u = divmod(r, 10)
        partes.append(_DECENAS[d] + (f" Y {_UNIDADES[u]}" if u else ""))
    return " ".join(partes)


def numero_a_letras(n: int) -> str:
    if n == 0:
        return "CERO"
    if n >= 1_000_000:
        mill, resto = divmod(n, 1_000_000)
        pre = "UN MILLON" if mill == 1 else f"{numero_a_letras(mill)} MILLONES"
        return f"{pre} {numero_a_letras(resto)}".strip() if resto else pre
    if n >= 1000:
        mil, resto = divmod(n, 1000)
        pre = "MIL" if mil == 1 else f"{_menor_mil(mil)} MIL"
        return f"{pre} {_menor_mil(resto)}".strip() if resto else pre
    return _menor_mil(n)


def importe_en_letras(importe: Decimal) -> str:
    entero = int(importe)
    cent = int((importe - entero) * 100)
    letras = numero_a_letras(entero)
    sufijo = "NUEVO SOL" if entero == 1 else "NUEVOS SOLES"
    return f"{letras} CON {cent:02d}/100 {sufijo}" if cent else f"{letras} {sufijo}"


def _reemplazar_en_parrafo(p, datos: dict) -> None:
    for run in p.runs:
        if "{{" in run.text:
            texto = run.text
            for k, v in datos.items():
                texto = texto.replace("{{" + k + "}}", v)
            run.text = texto


def _recorrer_parrafos(doc):
    for p in doc.paragraphs:
        yield p
    for tabla in doc.tables:
        for fila in tabla.rows:
            vistas = set()
            for celda in fila.cells:
                if id(celda._tc) in vistas:
                    continue
                vistas.add(id(celda._tc))
                for p in celda.paragraphs:
                    yield p


def _fmt_fecha(f) -> str:
    return f.strftime("%d / %m / %Y")


def _fmt_monto(v) -> str:
    return f"{Decimal(v):,.2f}" if v is not None else "-"


def construir_datos(sep, proyecto_nombre: str) -> dict:
    cli, lote = sep.cliente, sep.lote
    juridica = cli.tipo_persona == "juridica"
    nombre = cli.razon_social if juridica else f"{cli.nombres} {cli.apellidos}"
    documento = (cli.ruc if juridica else cli.numero_documento) or "-"
    return {
        "NUMERO": f"{sep.id:05d}",
        "FECHA": _fmt_fecha(sep.fecha_inicio),
        "IMPORTE": _fmt_monto(sep.importe),
        "FORMA_PAGO": (sep.medio_pago or (sep.tipo_pago.value if sep.tipo_pago else "-")).upper(),
        "VOUCHER": sep.voucher_numero or "-",
        "CLIENTE": nombre.upper(),
        "DNI": documento,
        "DIRECCION": cli.direccion or "",
        "EMAIL": cli.correo or "",
        "CELULAR1": cli.telefono or "",
        "CELULAR2": cli.segundo_contacto_telefono or "",
        "IMPORTE_LETRAS": importe_en_letras(Decimal(sep.importe)),
        "CONCEPTO": (sep.motivo or "SEPARACIÓN DE LOTE").upper().rstrip(".") + ".",
        "PARTIDA": lote.partida_registral or "-",
        "PROYECTO": f"PROYECTO INMOBILIARIO “{proyecto_nombre.upper()}”",
        "PRECIO_REGULAR": _fmt_monto(lote.precio_total_base),
        "PRECIO_DESCONTADO": _fmt_monto(
            sep.precio_descontado if sep.precio_descontado is not None else lote.precio_total_contado
        ),
        "MANZANA": (lote.manzana_nombre or "-").upper(),
        "LOTE": lote.codigo,
        "AREA": f"{Decimal(lote.area_m2):,.2f}",
        "VENCIMIENTO": _fmt_fecha(sep.fecha_vencimiento),
        "PRECIO_VALIDO_HASTA": _fmt_fecha(sep.precio_valido_hasta or sep.fecha_vencimiento),
    }


def generar_docx(sep, proyecto_nombre: str, con_clausulas: bool) -> bytes:
    doc = docx.Document(RUTA_PLANTILLA)
    datos = construir_datos(sep, proyecto_nombre)
    datos["TITULO"] = "CONTRATO" if con_clausulas else "PROFORMA"
    if con_clausulas:
        datos["ETIQUETA_FECHA"] = "Vencimiento de contrato de separación"
        datos["FECHA_LIMITE"] = datos["VENCIMIENTO"]
    else:
        datos["ETIQUETA_FECHA"] = "Precio válido hasta"
        datos["FECHA_LIMITE"] = datos["PRECIO_VALIDO_HASTA"]
    for p in _recorrer_parrafos(doc):
        _reemplazar_en_parrafo(p, datos)

    if not con_clausulas:
        tabla = doc.tables[0]
        # filas 7 y 8 = encabezado y texto de las cláusulas
        for fila in list(tabla.rows)[7:]:
            tabla._tbl.remove(fila._tr)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()



def _convertir_a_pdf(contenido_docx: bytes) -> bytes | None:
    """Convierte el .docx a PDF con LibreOffice. Devuelve None si no está disponible."""
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return None
    with tempfile.TemporaryDirectory() as tmp:
        origen = os.path.join(tmp, "doc.docx")
        with open(origen, "wb") as f:
            f.write(contenido_docx)
        try:
            subprocess.run(
                [exe, "--headless", "--convert-to", "pdf", "--outdir", tmp, origen],
                check=True, timeout=90, capture_output=True,
                env={**os.environ, "HOME": tmp},
            )
        except Exception:
            return None
        destino = os.path.join(tmp, "doc.pdf")
        if not os.path.isfile(destino):
            return None
        with open(destino, "rb") as f:
            return f.read()


def guardar_documento(prefijo: str, id_separacion: int, contenido: bytes, url_base: str) -> str:
    """Guarda el documento (PDF si se puede, si no .docx) y devuelve su URL pública."""
    pdf = _convertir_a_pdf(contenido)
    extension, datos = ("pdf", pdf) if pdf else ("docx", contenido)
    nombre = f"{prefijo}_{id_separacion}_{uuid.uuid4().hex[:8]}.{extension}"
    with open(os.path.join(DIR_DOCUMENTOS, nombre), "wb") as f:
        f.write(datos)
    return f"{url_base.rstrip('/')}/static/separaciones/{nombre}"